"""Consistent fail-closed tables and downloads; fitted curves, no new estimators."""
import io
import json
import zipfile
import numpy as np
import matplotlib.pyplot as plt
from reports.irt_b14_report import safe_frame,export_xlsx


def figure(tables,item):
    fig,axes=plt.subplots(2,2,figsize=(12,8),layout='constrained')
    p=tables['Curves'];d=tables['Difference'];I=tables['IIF'];T=tables['TIF']
    for (group,category),part in p[p.Item==item].groupby(['Group','Category']):
        axes[0,0].plot(part.Theta,part.Probability,label=f'{group}: {category}',linestyle='--' if group=='focal' else '-')
    for category,part in d[d.Item==item].groupby('Category'):
        axes[0,1].plot(part.Theta,part.Focal_minus_reference,label=str(category))
    for group,part in I[I.Item==item].groupby('Group'):axes[1,0].plot(part.Theta,part.Information,label=group)
    for group,part in T.groupby('Group'):axes[1,1].plot(part.Theta,part.Information,label=group)
    for ax,title in zip(axes.flat,(f'{item}: category probabilities','Focal − reference probability',f'{item}: item information','Test information')):
        ax.set(title=title,xlabel='Common theta');ax.legend(fontsize=7);ax.grid(alpha=.2)
    fig.suptitle('NOT CERTIFIED — descriptive multigroup ML')
    return fig


def png(tables,item):
    fig=figure(tables,item);buf=io.BytesIO();fig.savefig(buf,format='png',dpi=120);plt.close(fig)
    return buf.getvalue()


def narrative(result):
    m=result['metadata']
    return ('# Python mirt — DIF / NOT CERTIFIED\n\n'+m['Limitations']+'\n\n'
        'Yöntem: marjinal ML; açık çapalara dayalı add tasarımı. H1 seçilen aileleri tüm çapa dışı '
        'maddelerde serbest bırakır; her H0 bir incelenen maddeyi eşitler. Test edilmeyen aileler '
        'gruplar arasında eşittir. Rasch ve bu API’nin 1PL seçeneğinde a=1; uygulamanın ortak-serbest-a '
        '1PL kestirimiyle aynı değildir. Referans N(0,1); odak ortalama/varyansı çapa tanımlamasıyla '
        'kestirilir. Normal tek boyutlu gizil dağılım ve yerel bağımsızlık varsayılır.\n\n'
        'LR=2(logLik_H1−logLik_H0); negatif LR kırpılmaz. df serbest parametre sayısı farkıdır. '
        'AIC/BIC uyum ölçütleridir, mutlak model uygunluğu kanıtı değildir. SE/GA/Wald verilmez. '
        'Seçilen çoklu test düzeltmesi p değerleri kapalı olduğundan uygulanmamıştır; BH için bağımlılık '
        'varsayımları ve veriyle çapa seçiminin etkileri ayrıca doğrulanmalıdır.\n\n'
        'Etki: ortak [-4,4] gridinde en büyük mutlak kategori olasılık farkı; eşik/sınıflandırma yoktur. '
        'DTF: yalnız sıralı/ikili modellerde kategori koduyla beklenen toplam puanın odak−referans farkı; '
        'bir test istatistiği değildir. NRM için DTF ve Wright Map yoktur. Fark eğrileri gizil grup '
        'yoğunluklarıyla ağırlıklandırılmaz. Çapa duyarlılığı varsa ikinci açık çapa tasarımıdır; '
        'bu karşılaştırma çapa geçerliliğini kanıtlamaz.\n\n'
        'Tam girdiler yüklenmez; tekrar üretim için orijinal veriyle kayıtlı hash, eşleme ve ayarlar '
        'kullanılmalıdır. Yakınsamayan kestirimler yalnız geçici tanı çıktısıdır.\n\n```json\n'+
        json.dumps(m,ensure_ascii=False,indent=2)+'\n```\n')


def downloads(result,item):
    tables=result['tables']
    csv=safe_frame(tables['DIF']).to_csv(index=False).encode('utf-8-sig')
    xlsx=export_xlsx(tables)
    # Reuse the existing safe exporter; add familiar frozen, filtered blue headers.
    from openpyxl import load_workbook
    from openpyxl.styles import Font,PatternFill
    wb=load_workbook(io.BytesIO(xlsx))
    for ws in wb:
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='1E3A8A')
    buf=io.BytesIO();wb.save(buf);xlsx=buf.getvalue()
    picture=png(tables,item);report=narrative(result)
    buf=io.BytesIO()
    with zipfile.ZipFile(buf,'w',zipfile.ZIP_DEFLATED) as z:
        for name,frame in tables.items():z.writestr(name+'.csv',safe_frame(frame).to_csv(index=False))
        z.writestr('report.md',report);z.writestr('settings.json',json.dumps(result['metadata'],ensure_ascii=False,indent=2))
        z.writestr('report.xlsx',xlsx);z.writestr('curves.png',picture)
    return dict(csv=csv,xlsx=xlsx,png=picture,zip=buf.getvalue(),report=report.encode('utf-8'))
