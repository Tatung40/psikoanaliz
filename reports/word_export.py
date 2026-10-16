"""
PsikoAnaliz — Akademik Word (.docx) Raporlama Motoru (reports/word_export.py)
APA 7 uyumlu tablolar, metodolojik yorumlar, formül açıklamaları ve akademik değerlendirme raporu.
Geliştirici: Doç. Dr. Mehmet Taha Eser
"""

import io
from typing import Dict, List, Optional, Any
import numpy as np
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn


def set_cell_border(cell, **kwargs):
    """
    APA 7 formatında hücre kenarlıkları uygular (Yalnızca üst, alt ve başlık altı çizgileri).
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'<w:top w:val="{kwargs.get("top", "none")}" w:sz="{kwargs.get("top_sz", "4")}" w:space="0" w:color="{kwargs.get("top_color", "auto")}"/>\n'
        f'<w:left w:val="{kwargs.get("left", "none")}" w:sz="4" w:space="0" w:color="auto"/>\n'
        f'<w:bottom w:val="{kwargs.get("bottom", "none")}" w:sz="{kwargs.get("bottom_sz", "4")}" w:space="0" w:color="{kwargs.get("bottom_color", "auto")}"/>\n'
        f'<w:right w:val="{kwargs.get("right", "none")}" w:sz="4" w:space="0" w:color="auto"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)


def create_word_report(
    test_summary: Dict[str, Any],
    item_summary_df: pd.DataFrame,
    reliability_data: Dict[str, Any],
    distractor_data: Optional[Dict[str, Dict[str, Any]]] = None,
    group_data: Optional[Dict[str, Any]] = None,
    test_title: str = "Klasik Test Kuramı Psikometrik Değerlendirme Raporu"
) -> io.BytesIO:
    """
    APA 7 standardında akademik araştırma raporu üretir.
    """
    doc = Document()

    # Sayfa Kenar Boşlukları (Standart 1 inç / 2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Başlık
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title_p.add_run(test_title.upper())
    run_title.font.name = "Times New Roman"
    run_title.font.size = Pt(14)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(15, 23, 42)

    # Yazar ve Kurum
    author_p = doc.add_paragraph()
    author_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_author = author_p.add_run("Doç. Dr. Mehmet Taha Eser\nPsikoAnaliz — Klasik Test Kuramı Analiz Platformu")
    run_author.font.name = "Times New Roman"
    run_author.font.size = Pt(11)
    run_author.font.italic = True

    doc.add_paragraph()  # Boşluk

    # 1. GİRİŞ VE TEST METADATASI
    h1 = doc.add_heading(level=1)
    run_h1 = h1.add_run("1. Giriş ve Test Özellikleri")
    run_h1.font.name = "Times New Roman"
    run_h1.font.size = Pt(12)
    run_h1.font.bold = True
    run_h1.font.color.rgb = RGBColor(30, 58, 138)

    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.line_spacing = 1.15
    p_intro.paragraph_format.first_line_indent = Inches(0.5)
    run_intro = p_intro.add_run(
        f"Bu raporda sunulan analizler, Klasik Test Kuramı (KTK) prensiplerine uygun olarak gerçekleştirilmiştir. "
        f"Veri setinde toplam {test_summary.get('n_examinees', 0)} katılımcı ve {test_summary.get('k_items', 0)} madde "
        f"bulunmaktadır. Test puan ortalaması {test_summary.get('mean', 0.0):.2f} (SS = {test_summary.get('std_dev', 0.0):.2f}) "
        f"olarak hesaplanmıştır. Analiz kapsamında test düzeyi normlar, madde güçlük ve ayırt edicilik indeksleri, "
        f"güvenirlik katsayıları ve pedagojik tanısal bayraklar incelenmiştir."
    )
    run_intro.font.name = "Times New Roman"
    run_intro.font.size = Pt(11)

    # 2. TEST DÜZEYİ BETİMSEL İSTATİSTİKLER (TABLO 1)
    h2 = doc.add_heading(level=1)
    run_h2 = h2.add_run("2. Test Düzeyi Betimsel İstatistikler")
    run_h2.font.name = "Times New Roman"
    run_h2.font.size = Pt(12)
    run_h2.font.bold = True
    run_h2.font.color.rgb = RGBColor(30, 58, 138)

    # Tablo 1 Başlığı (APA 7)
    p_t1_title = doc.add_paragraph()
    r_t1_num = p_t1_title.add_run("Tablo 1\n")
    r_t1_num.bold = True
    r_t1_name = p_t1_title.add_run("Test Toplam Puanlarına İlişkin Betimsel ve Dağılım İstatistikleri")
    r_t1_name.italic = True
    for r in [r_t1_num, r_t1_name]:
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)

    table1 = doc.add_table(rows=1, cols=4)
    table1.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells1 = table1.rows[0].cells
    hdr_titles1 = ["İstatistik", "Değer", "İstatistik", "Değer"]
    for i, title in enumerate(hdr_titles1):
        hdr_cells1[i].text = title
        set_cell_border(hdr_cells1[i], top="single", top_sz="8", bottom="single", bottom_sz="8")
        p = hdr_cells1[i].paragraphs[0]
        p.runs[0].font.name = "Times New Roman"
        p.runs[0].font.size = Pt(10)
        p.runs[0].font.bold = True

    t1_data = [
        ("Katılımcı Sayısı (N)", str(test_summary.get("n_examinees", "-")), "Minimum Puan", str(test_summary.get("min", "-"))),
        ("Madde Sayısı (k)", str(test_summary.get("k_items", "-")), "Maksimum Puan", str(test_summary.get("max", "-"))),
        ("Ortalama (Mean)", str(test_summary.get("mean", "-")), "Ranj (Aralık)", str(test_summary.get("range", "-"))),
        ("Medyan (Ortanca)", str(test_summary.get("median", "-")), "Çarpıklık (Skewness)", str(test_summary.get("skewness", "-"))),
        ("Standart Sapma (s)", str(test_summary.get("std_dev", "-")), "Basıklık (Kurtosis)", str(test_summary.get("kurtosis", "-"))),
        ("Örneklem Varyansı (s²)", str(test_summary.get("variance", "-")), "Çeyrekler Arası Ranj (IQR)", str(test_summary.get("iqr", "-"))),
        ("Ort. Standart Hatası (SEM_m)", str(test_summary.get("sem_mean", "-")), "Taban / Tavan %", f"%{test_summary.get('floor_pct', 0.0):.1f} / %{test_summary.get('ceiling_pct', 0.0):.1f}")
    ]

    for row_idx, (s1, v1, s2, v2) in enumerate(t1_data):
        row_cells = table1.add_row().cells
        for col_idx, text in enumerate([s1, v1, s2, v2]):
            row_cells[col_idx].text = text
            p = row_cells[col_idx].paragraphs[0]
            p.runs[0].font.name = "Times New Roman"
            p.runs[0].font.size = Pt(10)
            is_last = (row_idx == len(t1_data) - 1)
            if is_last:
                set_cell_border(row_cells[col_idx], bottom="single", bottom_sz="8")

    # Tablo 1 Notu
    p_t1_note = doc.add_paragraph()
    r_note1 = p_t1_note.add_run("Not. Örneklem varyansı ve standart sapmasında serbestlik derecesi düzeltmesi (ddof = 1) kullanılmıştır.")
    r_note1.font.name = "Times New Roman"
    r_note1.font.size = Pt(9)
    r_note1.font.italic = True

    # 3. GÜVENİRLİK VE ÖLÇME HATASI (TABLO 2)
    h3 = doc.add_heading(level=1)
    run_h3 = h3.add_run("3. Güvenirlik ve Ölçme Hatası Bulguları")
    run_h3.font.name = "Times New Roman"
    run_h3.font.size = Pt(12)
    run_h3.font.bold = True
    run_h3.font.color.rgb = RGBColor(30, 58, 138)

    p_t2_title = doc.add_paragraph()
    r_t2_num = p_t2_title.add_run("Tablo 2\n")
    r_t2_num.bold = True
    r_t2_name = p_t2_title.add_run("Ölçme Aracına İlişkin Güvenirlik Katsayıları ve Ölçme Hatası Kestirimleri")
    r_t2_name.italic = True
    for r in [r_t2_num, r_t2_name]:
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)

    table2 = doc.add_table(rows=1, cols=3)
    table2.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells2 = table2.rows[0].cells
    for i, title in enumerate(["Güvenirlik Katsayısı / Metrik", "Kestirilen Değer", "Açıklama / Değerlendirme"]):
        hdr_cells2[i].text = title
        set_cell_border(hdr_cells2[i], top="single", top_sz="8", bottom="single", bottom_sz="8")
        p = hdr_cells2[i].paragraphs[0]
        p.runs[0].font.name = "Times New Roman"
        p.runs[0].font.size = Pt(10)
        p.runs[0].font.bold = True

    t2_data = [
        ("Cronbach Alfa (α)", str(reliability_data.get("cronbach_alpha", "-")), "İç tutarlık katsayısı"),
        ("Standartlaştırılmış Alfa", str(reliability_data.get("standardized_alpha", "-")), "Maddeler arası ortalama korelasyona dayalı"),
        ("KR-20 (Kuder-Richardson)", str(reliability_data.get("kr20", "-")) if reliability_data.get("kr20") is not None else "Çok Kategorili Ölçek", "İkili veride Alfa ile tam eşdeğer"),
        ("KR-21 Katsayısı", str(reliability_data.get("kr21", "-")) if reliability_data.get("kr21") is not None else "-", "Eşit güçlük varsayımı altında alt sınır"),
        ("İki-Yarı (Tek-Çift) Spearman-Brown", str(reliability_data.get("spearman_brown_odd_even", "-")), "Test yarısı eşdeğerlik kestirimi"),
        ("Guttman Split-Half (λ4)", str(reliability_data.get("guttman_lambda4", "-")), "Yarı-test varyans temelli güvenirlik"),
        ("McDonald Omega (ω_t)", str(reliability_data.get("mcdonald_omega", "-")), "Tek faktörlü model faktör yüklerine dayalı"),
        ("Ölçmenin Standart Hatası (SEM)", str(reliability_data.get("sem", "-")), f"Bireysel puan hata marjı: ±{reliability_data.get('ci_95_margin', '-')} (%95 GA)"),
        ("Feldt (1969) %95 Alfa GA", f"[{reliability_data.get('alpha_ci_lower', '-')}, {reliability_data.get('alpha_ci_upper', '-')}]", "Kuramsal güven aralığı")
    ]

    for row_idx, (m, val, desc) in enumerate(t2_data):
        row_cells = table2.add_row().cells
        for col_idx, text in enumerate([m, val, desc]):
            row_cells[col_idx].text = text
            p = row_cells[col_idx].paragraphs[0]
            p.runs[0].font.name = "Times New Roman"
            p.runs[0].font.size = Pt(10)
            if row_idx == len(t2_data) - 1:
                set_cell_border(row_cells[col_idx], bottom="single", bottom_sz="8")

    # Tablo 2 Notu
    p_t2_note = doc.add_paragraph()
    r_note2 = p_t2_note.add_run(f"Not. Değerlendirme: {reliability_data.get('reliability_evaluation', 'Hesaplandı.')}")
    r_note2.font.name = "Times New Roman"
    r_note2.font.size = Pt(9)
    r_note2.font.italic = True

    # 4. MADDE ANALİZİ BULGULARI (TABLO 3)
    h4 = doc.add_heading(level=1)
    run_h4 = h4.add_run("4. Klasik Madde Analizi ve Tanısal Değerlendirmeler")
    run_h4.font.name = "Times New Roman"
    run_h4.font.size = Pt(12)
    run_h4.font.bold = True
    run_h4.font.color.rgb = RGBColor(30, 58, 138)

    p_t3_title = doc.add_paragraph()
    r_t3_num = p_t3_title.add_run("Tablo 3\n")
    r_t3_num.bold = True
    r_t3_name = p_t3_title.add_run("Maddelerin Klasik Güçlük, Ayırt Edicilik ve Güvenirlik İndeksleri")
    r_t3_name.italic = True
    for r in [r_t3_num, r_t3_name]:
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)

    t3_cols = ["Madde", "N", "Güçlük (p)", "r_it", "r_ir", "D (%27)", "Sil. Alfa", "Pedagojik Değerlendirme"]
    table3 = doc.add_table(rows=1, cols=len(t3_cols))
    table3.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, title in enumerate(t3_cols):
        cell = table3.rows[0].cells[i]
        cell.text = title
        set_cell_border(cell, top="single", top_sz="8", bottom="single", bottom_sz="8")
        p = cell.paragraphs[0]
        p.runs[0].font.name = "Times New Roman"
        p.runs[0].font.size = Pt(9)
        p.runs[0].font.bold = True

    for row_idx, (_, row) in enumerate(item_summary_df.iterrows()):
        row_cells = table3.add_row().cells
        vals = [
            str(row["Madde"]),
            str(row["N"]),
            f"{row['Güçlük (p)']:.2f}" if not pd.isna(row["Güçlük (p)"]) else "-",
            f"{row['r_it']:.2f}" if not pd.isna(row["r_it"]) else "-",
            f"{row['r_ir (Düzeltilmiş)']:.2f}" if not pd.isna(row["r_ir (Düzeltilmiş)"]) else "-",
            f"{row['D (%27)']:.2f}" if not pd.isna(row["D (%27)"]) else "-",
            f"{row['Silindiğinde Alfa']:.2f}" if not pd.isna(row["Silindiğinde Alfa"]) else "-",
            str(row["Tanısal Bayrak"])
        ]
        for col_idx, val_str in enumerate(vals):
            cell = row_cells[col_idx]
            cell.text = val_str
            p = cell.paragraphs[0]
            p.runs[0].font.name = "Times New Roman"
            p.runs[0].font.size = Pt(9)
            if row_idx == len(item_summary_df) - 1:
                set_cell_border(cell, bottom="single", bottom_sz="8")

    # 5. YÖNTEMSEL KAYNAKÇA (APA 7)
    h5 = doc.add_heading(level=1)
    run_h5 = h5.add_run("5. Kaynakça")
    run_h5.font.name = "Times New Roman"
    run_h5.font.size = Pt(12)
    run_h5.font.bold = True
    run_h5.font.color.rgb = RGBColor(30, 58, 138)

    refs = [
        "Crocker, L., & Algina, J. (1986). Introduction to classical and modern test theory. Holt, Rinehart and Winston.",
        "Feldt, L. S. (1969). A test of the hypothesis that Cronbach's alpha or Kuder-Richardson coefficient twenty is the same for two tests. Psychometrika, 34(3), 363-373.",
        "Kelley, T. L. (1939). The selection of upper and lower groups for the validation of test items. Journal of Educational Psychology, 30(1), 17-24.",
        "Lord, F. M., & Novick, M. R. (1968). Statistical theories of mental test scores. Addison-Wesley."
    ]

    for ref in refs:
        p_ref = doc.add_paragraph()
        p_ref.paragraph_format.line_spacing = 1.15
        p_ref.paragraph_format.left_indent = Inches(0.5)
        p_ref.paragraph_format.first_line_indent = Inches(-0.5)  # Hanging indent
        r_ref = p_ref.add_run(ref)
        r_ref.font.name = "Times New Roman"
        r_ref.font.size = Pt(10)

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output
