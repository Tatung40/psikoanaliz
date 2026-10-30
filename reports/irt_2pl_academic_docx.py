"""Read-only 2PL Word report using fitted Python mirt LibraryReport.

No new calibration, API access, or unverified inferential statistics.
"""
from copy import deepcopy
from io import BytesIO
import math

from docx import Document
from docx.shared import Inches
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

from core.academic_result_adapters import irt_report_to_result
from core.irt_python_mirt_report import build_report
from reports.academic_docx import build_academic_docx


def _png(fig):
    buffer = BytesIO()
    fig.tight_layout()
    fig.savefig(buffer, format="png", dpi=150)
    plt.close(fig)
    buffer.seek(0)
    return buffer


def export_2pl_docx(outcome, items):
    """Produce DOCX bytes only for a converged, complete 2PL fit."""
    report = build_report(outcome, model="2PL", items=items)
    result = irt_report_to_result(report)
    # The generic manifest cannot distinguish item/category series: suppress its
    # plotting and use explicitly grouped ICC and item-information plots below.
    document_source = deepcopy(result)
    document_source["metadata"]["analysis_type"] = "2pl_word_grouped"
    doc = Document(BytesIO(build_academic_docx(document_source)))
    curves = report.curves
    information = report.information
    required = {"Madde", "Theta", "Kategori", "Olasılık"}
    if not required.issubset(curves.columns):
        raise ValueError("2PL curves lack grouping coordinates")
    if "θ" not in information.columns:
        raise ValueError("2PL information lacks theta axis")
    doc.add_heading("Madde karakteristik eğrileri (ICC)", level=1)
    for item in items:
        selected = curves[curves["Madde"].astype(str) == str(item)]
        if selected.empty:
            raise ValueError(f"Missing ICC for item {item}")
        fig, ax = plt.subplots(figsize=(6, 3.4))
        for category, part in selected.groupby("Kategori", sort=True):
            part = part.sort_values("Theta")
            ax.plot(part["Theta"], part["Olasılık"], label=f"Kategori {category}")
        ax.set(xlabel="θ", ylabel="Yanıt olasılığı", title=f"Madde: {item}", ylim=(0, 1))
        ax.legend()
        doc.add_paragraph(f"Şekil: {item} ICC")
        doc.add_picture(_png(fig), width=Inches(5.7))
    doc.add_heading("Madde ve test bilgi fonksiyonları", level=1)
    fig, ax = plt.subplots(figsize=(6, 3.5))
    ordered = information.sort_values("θ")
    for item in items:
        key = f"I_{item}"
        if key not in ordered:
            plt.close(fig)
            raise ValueError(f"Missing IIF column {key}")
        ax.plot(ordered["θ"], ordered[key], label=str(item), alpha=0.65)
    if "Test bilgisi" in ordered:
        ax.plot(ordered["θ"], ordered["Test bilgisi"], label="Test bilgisi", linewidth=2)
    ax.set(xlabel="θ", ylabel="Bilgi", title="IIF ve TIF")
    ax.legend(fontsize=7, ncol=2)
    doc.add_picture(_png(fig), width=Inches(5.7))
    doc.add_paragraph("Not. Grafikler doğrudan Python mirt kalibrasyonundan hesaplanan eğri ve bilgi tablolarını kullanır. Standart hata ve çıkarımsal karar üretilmez.")
    output = BytesIO()
    doc.save(output)
    return output.getvalue()
