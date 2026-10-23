"""Qualified Turkish categorical DIF bundles, without calibration or person records."""
import io
import importlib.metadata
import sys
import zipfile
import numpy as np
import pandas as pd
from core import polytomous_dif as P
from reports.irt_b14_report import safe_frame, export_xlsx

REFERENCES = {
    "PCM": "Masters, G. N. (1982). A Rasch model for partial credit scoring. Psychometrika, 47(2), 149–174. https://doi.org/10.1007/BF02296272",
    "GPCM": "Muraki, E. (1992). A generalized partial credit model: Application of an EM algorithm. Applied Psychological Measurement, 16(2), 159–176. https://doi.org/10.1177/014662169201600206",
    "RSM": "Andrich, D. (1978). A rating formulation for ordered response categories. Psychometrika, 43(4), 561–573. https://doi.org/10.1007/BF02293814",
    "NRM": "Bock, R. D. (1972). Estimating item parameters and latent ability when responses are scored in two or more nominal categories. Psychometrika, 37(1), 29–51. https://doi.org/10.1007/BF02291411",
}
COMMON_CAUTION = "Bağımsız sayısal doğrulama ve benzetimler sınırlı tasarımlardadır; üretim/genel FDR sertifikası değildir. R mirt bu çalışmada çalıştırılmadı. Eksik yanıtlar MAR altında atlanır; tüm iç içe modeller aynı kişi/yanıt kodlarını kullanır. Kategori marjinleri betimseldir; tek boyutluluk/yerel bağımsızlığı doğrulamaz."
MODEL_CAUTION = {
    "PCM": "Bu PCM a=1 yanında referans varyansını 1 kabul eden kısıtlı modeldir. Bu ek dağılım varsayımı serbest referans varyanslı Rasch kalibrasyonu değildir; grupları ters çevirmek aynı modelin salt yeniden ölçeklemesi olmayabilir. Adımlar sıralanmaz.",
    "GPCM": "Adımlar theta birimindedir; eğim değişirken adım eşitliği kanonik kesişim eşitliği değildir. Genel, koşullu adım ve koşullu eğim ayrı hipotezlerdir.",
    "RSM": "a=1 ve referans varyansı 1 ek dağılım varsayımıdır; grup terslemesi aynı modelin salt yeniden ölçeklemesi olmayabilir. Ortak adımlar tüm maddeler/gruplarda aynıdır; yalnız konum DIF'i sınanır. Ayrı R mirt çok gruplu kalibrasyonu yoktur.",
    "NRM": "Kategori kodları sırasızdır; ilk kategori referanstır. İlk çapanın ikinci/referans kategori eğim farkı eksen işaretini belirler. Parametre sınırına yaklaşan alternatif kodlama da çıkarımı engeller. Puan veya başarı/avantaj yorumu yapılmaz. Ağırlıklı kategori-etki aralıklarında benzetim kapsamı %88'e düştü; nominal %95 kapsam doğrulanmadığından bu etki aralıkları raporda/arayüzde verilmez. Noktasal eğri/parametre aralıkları da deneysel normal yaklaşımıdır.",
}


def report_tables(out):
    inputs = out["inputs"]; full = next(iter(out["fits"].values()))["H2"]
    d, p = full["design"], full["params"]
    A, D, mu, sd = d.unpack(p)
    meta = {"Model": out["model"], "Bilimsel durum": out["status"], "Yöntem": out["method"],
        "Referans": str(out["reference"]), "Odak": str(out["focal"]), "N referans": inputs["sizes"][0],
        "N odak": inputs["sizes"][1], "Dışlanan satır": inputs["excluded_rows"], "Madde sayısı": len(inputs["names"]),
        "Kategori etiketleri (ilk etiket referans)": repr(inputs["category_labels"]),
        "Çapalar": ", ".join(map(str, out["anchors"])), "Çapa arındırma": "Yok; önceden belirlenmiş çapalar",
        "Odak ortalama (impact)": mu, "Odak SS (impact)": sd, "BH nominal alfa": out["alpha"],
        "Quadrature": out["points"], "Denetim quadrature": max(out["points"] + 20, 81),
        "Uyarlamalı integral": "Denetim başarısızsa yalnız ilgili fit +20 noktayla yeniden kestirilir; en çok 161, denetim en çok 181",
        "Quadrature logL toleransı": .001, "Skor/kişi toleransı": 2e-6,
        "En az kategori hücresi": inputs["min_cell"], "Optimizasyon sayısı": out["n_fits"],
        "Parametre sınırları": repr(dict(zip(map(str, d.labels), d.bounds[:d.latent]))),
        "Gizil sınırlar": "Odak mu [-3,3], SS [.3,3]; referans N(0,1)",
        "Tanımlama": "eta_c=A_c*theta+D_c; eta_0=0. PCM/RSM a=1. RSM ortak adım toplamı sıfır. NRM yönü ilk çapanın ikinci/referans kategori eğim farkıyla tanımlı.",
        "Tekrarlanabilirlik": "Deterministik L-BFGS-B, analitik skor; eğriler/rapor için tekrar kestirim yok",
        "Python": sys.version.split()[0], **{q: importlib.metadata.version(q) for q in ("numpy", "scipy", "pandas")}}
    diagnostics = []
    for i, fits in out["fits"].items():
        for h, fit in fits.items():
            npar = fit["design"].size
            diagnostics.append({"Madde": inputs["names"][i], "Hipotez": h, "logL": fit["loglik"], "Serbest parametre": npar,
                "AIC": 2 * npar - 2 * fit["loglik"], "BIC": npar * np.log(sum(inputs["sizes"])) - 2 * fit["loglik"],
                "Yakınsama": fit["converged"], "Sınırda": fit["at_bound"], "Skor/kişi": fit["score"], "Yineleme": fit["iterations"],
                "Bilgi matrisi": fit["cov"] is not None, "En küçük bilgi özdeğeri": fit["information_min_eigenvalue"],
                "Bilgi koşul sayısı": fit["information_condition"], "Quadrature logL farkı": fit["quadrature_error"],
                "Quadrature noktası": fit["points"], "Denetim noktası": fit["audit_points"],
                "Quadrature doğrulandı": fit["quadrature_ok"], "Not": fit["reason"]})
    margins, coefficients, physical = [], [], []
    a, b, tau = d.physical(p)
    z, lw = P._grid(max(out["points"], 81))
    for g, label in ((0, str(out["reference"])), (1, str(out["focal"]))):
        predicted = np.einsum("q,qic->ic", np.exp(lw), np.exp(P.log_probabilities(z if g == 0 else mu + sd * z, A[g], D[g])))
        for i, name in enumerate(inputs["names"]):
            physical.append({"Grup": label, "Madde": name, "a veya kategori eğimleri": repr(a[g, i].tolist()),
                             "b veya kategori kesişimleri": repr(b[g, i].tolist()),
                             "Ortak RSM adımları": repr(tau.tolist()) if out["model"] == "RSM" else ""})
            observed = inputs["x"][inputs["groups"] == g, i]; n = (observed >= 0).sum()
            for c, category in enumerate(inputs["category_labels"]):
                coefficients.append({"Grup": label, "Madde": name, "Kategori": str(category), "A": A[g, i, c], "D": D[g, i, c]})
                count = int((observed == c).sum())
                margins.append({"Grup": label, "Madde": name, "Kategori": str(category), "Gözlenen n": count,
                    "Gözlenen oran": count / n, "Model oranı": predicted[i, c], "Oran farkı": count / n - predicted[i, c],
                    "Eksik n": int((observed < 0).sum())})
    categories, scores = [], []
    for i, models in out["fits"].items():
        category, score = P.conditional_curves(models["H2"], i, usable=np.isfinite(out["table"].loc[i, "p"]))
        category["Kategori"] = category.Kategori.map(dict(enumerate(inputs["category_labels"]))).astype(str)
        category.insert(0, "Madde", inputs["names"][i]); categories.append(category)
        if not score.empty: score.insert(0, "Madde", inputs["names"][i]); scores.append(score)
    effects, parameters = out["effects"].copy(), out["parameters"].copy()
    if out["model"] == "NRM":
        effects[["%95 GA alt", "%95 GA üst"]] = np.nan
        effects["Aralık durumu"] = "Kullanılamıyor: ağırlıklı kategori etkisinde benzetim kapsamı %88; genel %95 kapsam doğrulanmadı"
        effects["Etki"] = effects.Etki.map({f"Kategori {c} olasılık farkı": f"{label!s} olasılık farkı" for c, label in enumerate(inputs["category_labels"])})
        mapping = {f"Kategori {c} {kind}": f"{inputs['category_labels'][c]!s} / {inputs['category_labels'][0]!s} {kind}"
                   for c in range(1, inputs["categories"]) for kind in ("eğim", "kesişim")}
        parameters["Parametre"] = parameters.Parametre.map(mapping)
    result = dict(Yontem=pd.DataFrame({"Özellik": list(meta), "Değer": list(meta.values())}), Maddeler=out["table"].copy(),
        Parametreler=parameters, Etkiler=effects, Katsayilar=pd.DataFrame(coefficients), Fiziksel_parametreler=pd.DataFrame(physical),
        Optimizasyon=pd.DataFrame(diagnostics), Kategori_tanilari=pd.DataFrame(margins),
        Kategori_egrileri=pd.concat(categories, ignore_index=True))
    if scores: result["Puan_farki"] = pd.concat(scores, ignore_index=True)
    return result


def png_bytes(out, tables, item):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    nominal = out["model"] == "NRM"
    fig, axes = plt.subplots(2 if nominal else 3, 1, figsize=(7, 8), sharex=True)
    for j, (c, rows) in enumerate(tables["Kategori_egrileri"].loc[lambda f: f.Madde == item].groupby("Kategori", sort=False)):
        color = f"C{j}"
        axes[0].plot(rows.Theta, rows["P (referans)"], color=color, label=f"{out['reference']}, {c}")
        axes[0].plot(rows.Theta, rows["P (odak)"], color=color, ls="--", label=f"{out['focal']}, {c}")
        axes[1].plot(rows.Theta, rows["Olasılık farkı"], color=color, label=str(c))
        axes[1].fill_between(rows.Theta, rows["%95 GA alt"], rows["%95 GA üst"], color=color, alpha=.1)
    axes[0].set(ylabel="Kategori olasılığı", ylim=(0, 1), title=f"{item}: {out['model']}, ortak theta")
    axes[0].legend(fontsize=7, ncol=2)
    axes[1].set(ylabel="Odak − referans olasılık"); axes[1].axhline(0, color="gray", lw=.7)
    if not nominal:
        score = tables["Puan_farki"].loc[lambda f: f.Madde == item]
        axes[2].plot(score.Theta, score["Beklenen puan farkı"], color="black")
        axes[2].fill_between(score.Theta, score["%95 GA alt"], score["%95 GA üst"], alpha=.15)
        axes[2].set(ylabel="Beklenen puan farkı"); axes[2].axhline(0, color="gray", lw=.7)
    axes[-1].set(xlabel="Theta (referans N(0,1); NRM yönü beyan edilen çapa çiftine bağlı)")
    fig.text(.02, .005, "Bantlar noktasal %95 delta aralıklarıdır; NRM kategorileri başarı sırası değildir.", fontsize=7)
    fig.tight_layout(rect=(0, .02, 1, 1))
    buffer = io.BytesIO(); fig.savefig(buffer, format="png", dpi=300); plt.close(fig)
    return buffer.getvalue()


def bundle(out, tables=None):
    from docx import Document
    from docx.shared import Inches
    tables = report_tables(out) if tables is None else tables
    model = out["model"]; cautions = [out["status"], COMMON_CAUTION, MODEL_CAUTION[model]] + out["warnings"]
    reference = REFERENCES.get(model, "NRM: eta_c=A_c*theta+D_c; referans kategori eta_0=0. Kaynak/bağımsız doğrulama ayrıntıları POLYTOMOUS_DIF_VALIDATION.md dosyasındadır.")
    doc = Document(); doc.add_heading(f"{model}: sabit çapalı formel DIF araştırma raporu", 0)
    for note in cautions: doc.add_paragraph(note)
    for name in ("Yontem", "Maddeler", "Parametreler", "Etkiler"):
        frame = tables[name]
        if name == "Maddeler": frame = frame.reindex(columns=["Madde", "p", "q (BH)", "Karar"])
        doc.add_heading(name, 1)
        table = doc.add_table(rows=1, cols=len(frame.columns)); table.style = "Light Shading Accent 1"
        for cell, label in zip(table.rows[0].cells, frame.columns): cell.text = str(label)
        for row in frame.itertuples(index=False, name=None):
            for cell, value in zip(table.add_row().cells, row): cell.text = f"{value:.5g}" if isinstance(value, (float, np.floating)) else str(value)
    doc.add_paragraph("Tüm tanılar ve noktasal eğri aralıkları XLSX/CSV eklerindedir. Pozitif olasılık farkı yalnızca belirtilen kategorinin aynı theta'da odak grupta daha olası olduğunu gösterir. Sıralı modellerde pozitif puan etkisi daha yüksek beklenen puandır; NRM için puan/başarı/avantaj yorumu yapılmaz.")
    figures = [png_bytes(out, tables, item) for item in out["table"].loc[~out["table"]["Sabit çapa"], "Madde"]]
    for figure in figures: doc.add_picture(io.BytesIO(figure), width=Inches(6))
    doc.add_heading("Kaynak", 1); doc.add_paragraph(reference)
    word = io.BytesIO(); doc.save(word); output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(f"{model}_DIF.docx", word.getvalue())
        archive.writestr(f"{model}_DIF.xlsx", export_xlsx(tables))
        for name, frame in tables.items(): archive.writestr(f"{name}.csv", safe_frame(frame).to_csv(index=False).encode("utf-8-sig"))
        for j, figure in enumerate(figures): archive.writestr(f"item_{j + 1}_300dpi.png", figure)
        archive.writestr("README.txt", ("\n\n".join(cautions + [reference]) + "\nKişi kimlikleri/ham yanıtlar yoktur. CSV: UTF-8 BOM; grafik: 300 dpi. Rapor için yeniden kalibrasyon yapılmaz.").encode("utf-8"))
    return output.getvalue()
