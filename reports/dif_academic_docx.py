"""Offline Word export for a completed Python mirt DIF analysis.

Preserves the original four-panel DIF figure and all source tables.
"""
from io import BytesIO

from docx import Document
from docx.shared import Inches

from core.academic_result_adapters import dif_report_to_result
from reports.academic_docx import build_academic_docx
from reports.python_mirt_dif_report import png


def dif_academic_docx(result, item):
    tables = result["tables"]
    if item not in set(tables["Parameters"]["Item"]):
        raise ValueError("Selected item is absent from fitted parameters")
    envelope = dif_report_to_result(result)
    raw = build_academic_docx(envelope)
    document = Document(BytesIO(raw))
    document.add_heading("DIF — özgün dört panelli madde grafiği", level=1)
    document.add_paragraph(f"Madde: {item}. Bu görsel, mevcut çok gruplu Python mirt sonuçlarından üretilmiştir.")
    document.add_picture(BytesIO(png(tables, item)), width=Inches(6.0))
    document.add_paragraph("Not. Betimsel karşılaştırma; p/q ve istatistiksel DIF kararı kapalıdır.")
    output = BytesIO()
    document.save(output)
    return output.getvalue()
