"""Qualified Turkish GRM research bundle; no recalibration or respondent export."""
import io
import importlib.metadata
import sys
import zipfile
import numpy as np
import pandas as pd
from core import grm_dif as G
from reports.irt_b14_report import safe_frame, export_xlsx

REFERENCES = [
    "Samejima, F. (1968). Estimation of latent ability using a response pattern of graded scores "
    "(Research Bulletin RB-68-02). Educational Testing Service. https://doi.org/10.1002/j.2333-8504.1968.tb00153.x",
    "Chalmers, R. P. (2012). mirt: A multidimensional item response theory package for the R environment. "
    "Journal of Statistical Software, 48(6), 1–29. https://doi.org/10.18637/jss.v048.i06",
]

CAUTIONS = [
    "Deneysel kapsam: ortak 3–10 sıralı kategori, iki grup ve önceden belirlenmiş sabit çapalar. "
    "NRM/nominal seçenekler veya veriyle çapa seçimi bu yöntemle analiz edilmez.",
    "Bağımsız sayısal kontroller: analitik olasılıklar, uyarlamalı integral ve Python mirt 1.2.0 EM. "
    "Bağımsız kestirim ve benzetim örnekleri 3 kategorilidir. R mirt bu doğrulamada çalıştırılmadı; "
    "Chalmers kaynağı Python paketini tanımlamaz.",
    "Referans yetenek dağılımı N(0,1); odak ortalama ve standart sapması serbesttir. "
    "Yetenek dağılımı farkı (impact), DIF değildir. Eğriler aynı theta düzeyinde karşılaştırılır.",
    "H0: a,b ortak; Hb: b serbest/a ortak; Ha: a serbest/b ortak; H2: a,b serbest. "
    "Genel H2/H0 testi C sd; eşik H2/Ha testi C−1 sd; eğim H2/Hb testi 1 sd.",
    "Eşikler theta birimindeki b parametreleridir. d=−a*b olduğundan eğim serbestken b eşitliği, "
    "R mirt eğim-kesişim gösterimindeki d eşitliğiyle aynı hipotez değildir.",
    "Genel/eşik/eğim testleri ayrı BH aileleridir. Kullanılamayan p değerleri düzeltilmez; "
    "aileler arasında ortak FDR veya veriyle seçilmiş çapalar için hata kontrolü garanti edilmez.",
    "Noktasal Wald/delta yöntemi %95 aralıkları eşzamanlı değildir. Anlamlı fark madde yanlılığı "
    "kanıtı, anlamsız fark ise değişmezlik/eşdeğerlik kanıtı değildir. Madde otomatik silinmez.",
    "Tek boyutluluk ve koşullu yerel bağımsızlık varsayılır. Kategori marjin tanıları betimseldir; "
    "bu varsayımları doğrulayan formel model uygunluğu testi değildir.",
    "Eksik yanıtlar olabilirlikte atlanır (MAR varsayımı). Her model aynı kişi/kategori alt kümesini kullanır. "
    "Alt örneklem/hücre sınırları genel yeterlilik teoremi değildir. Gözlenen marjinler eksiklik mekanizmasına duyarlıdır.",
]


def report_tables(out):
    inputs = out["inputs"]
    full = next(iter(out["fits"].values()))["H2"]
    _, _, mu, sd = full["design"].unpack(full["params"])
    meta = {"Model": "GRM", "Bilimsel durum": out["status"], "Yöntem": out["method"],
        "Referans": str(out["reference"]), "Odak": str(out["focal"]), "N referans": inputs["sizes"][0],
        "N odak": inputs["sizes"][1], "Dışlanan satır": inputs["excluded_rows"],
        "Madde sayısı": len(inputs["names"]), "Kategori sayısı": inputs["categories"],
        "Çapalar": ", ".join(map(str, out["anchors"])), "Çapa arındırma": "Yok; sabit/değişmez kabul edilen çapalar",
        "Odak ortalama (impact)": mu, "Odak SS (impact)": sd, "BH nominal alfa": out["alpha"],
        "Quadrature": out["points"], "Denetim quadrature": max(out["points"] + 20, 81),
        "Quadrature logL toleransı": .001, "Denetim skor toleransı": 1e-5,
        "En az kategori hücresi": inputs["min_cell"], "Optimizasyon sayısı": out["n_fits"],
        "Sınırlar": "a: [.1,4]; ilk eşik: [-6,4]; ardışık eşik aralığı: [.05,6]; odak mu: [-3,3], SS: [.3,3]",
        "Tekrarlanabilirlik": "Deterministik L-BFGS-B; analitik skor; veriyle rastgele çapa seçimi yok",
        "Python": sys.version.split()[0],
        **{p: importlib.metadata.version(p) for p in ("numpy", "scipy", "pandas")}}
    diagnostics = []
    for i, models in out["fits"].items():
        for hypothesis, fit in models.items():
            parameters = fit["design"].size
            diagnostics.append({"Madde": inputs["names"][i], "Hipotez": hypothesis,
                "logL": fit["loglik"], "Serbest parametre": parameters,
                "AIC": 2 * parameters - 2 * fit["loglik"],
                "BIC": parameters * np.log(sum(inputs["sizes"])) - 2 * fit["loglik"],
                "Yakınsama": fit["converged"], "Sınırda": fit["at_bound"], "Skor/kişi": fit["score"],
                "Yineleme": fit["iterations"], "Bilgi matrisi": fit["cov"] is not None,
                "En küçük bilgi özdeğeri": fit["information_min_eigenvalue"],
                "Bilgi koşul sayısı": fit["information_condition"],
                "Quadrature logL farkı": fit["quadrature_error"], "Quadrature doğrulandı": fit["quadrature_ok"],
                "Not": fit["reason"]})
    margins, coefficients = [], []
    a, b, _, _ = full["design"].unpack(full["params"])
    z, lw = G._grid(max(out["points"], 61))
    for g, label in ((0, str(out["reference"])), (1, str(out["focal"]))):
        theta = z if g == 0 else mu + sd * z
        probabilities = np.exp(G.category_terms(theta, a[g], b[g])[0])
        predicted = np.einsum("q,qic->ic", np.exp(lw), probabilities)
        for i, name in enumerate(inputs["names"]):
            coefficients.append({"Grup": label, "Madde": name, "a": a[g, i],
                **{f"b{c + 1}": value for c, value in enumerate(b[g, i])}})
            observed = inputs["x"][inputs["groups"] == g, i]
            n = (observed >= 0).sum()
            for c in range(inputs["categories"]):
                count = int((observed == c).sum())
                margins.append({"Grup": label, "Madde": name, "Kategori": c, "Gözlenen n": count,
                    "Gözlenen oran": count / n, "Model oranı": predicted[i, c],
                    "Oran farkı": count / n - predicted[i, c], "Eksik n": int((observed < 0).sum())})
    categories, scores = [], []
    for i, models in out["fits"].items():
        usable = np.isfinite(out["table"].loc[i, "p"])
        category, score = G.conditional_curves(models["H2"], i, usable=usable)
        category.insert(0, "Madde", inputs["names"][i]); score.insert(0, "Madde", inputs["names"][i])
        categories.append(category); scores.append(score)
    return {"Yontem": pd.DataFrame({"Özellik": list(meta), "Değer": list(meta.values())}),
            "Maddeler": out["table"].copy(), "Parametreler": out["parameters"].copy(),
            "Katsayilar": pd.DataFrame(coefficients),
            "Optimizasyon": pd.DataFrame(diagnostics), "Kategori_tanilari": pd.DataFrame(margins),
            "Kategori_egrileri": pd.concat(categories, ignore_index=True),
            "Puan_farki": pd.concat(scores, ignore_index=True)}


def png_bytes(category, score, item, reference, focal):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    fig, axes = plt.subplots(3, 1, figsize=(7, 8), sharex=True)
    for c, rows in category.groupby("Kategori"):
        color = f"C{c}"
        axes[0].plot(rows.Theta, rows["P (referans)"], color=color, label=f"{reference}, kategori {c}")
        axes[0].plot(rows.Theta, rows["P (odak)"], color=color, ls="--", label=f"{focal}, kategori {c}")
        axes[1].plot(rows.Theta, rows["Olasılık farkı"], color=color, label=f"Kategori {c}")
        axes[1].fill_between(rows.Theta, rows["%95 GA alt"], rows["%95 GA üst"], color=color, alpha=.09)
    axes[0].set(ylabel="Kategori olasılığı", ylim=(0, 1), title=f"{item}: GRM, ortak theta ölçeği")
    axes[0].legend(fontsize=7, ncol=2)
    axes[1].axhline(0, color="gray", lw=.7); axes[1].set(ylabel="Odak − referans")
    axes[2].plot(score.Theta, score["Beklenen puan farkı"], color="black")
    axes[2].fill_between(score.Theta, score["%95 GA alt"], score["%95 GA üst"], alpha=.15, color="black")
    axes[2].axhline(0, color="gray", lw=.7)
    axes[2].set(xlabel="Theta (referans: N(0,1))", ylabel="Beklenen puan farkı")
    fig.text(.02, .005, "Bantlar noktasal %95 delta aralıklarıdır; model/çapa varsayımları koşulundadır.", fontsize=7)
    fig.tight_layout(rect=(0, .02, 1, 1))
    buffer = io.BytesIO(); fig.savefig(buffer, format="png", dpi=300); plt.close(fig)
    return buffer.getvalue()


def word_bytes(out, tables, figures):
    from docx import Document
    from docx.shared import Inches
    doc = Document()
    doc.add_heading("GRM: sabit çapalı formel DIF araştırma raporu", 0)
    doc.add_paragraph(out["status"])
    for note in CAUTIONS + out["warnings"]:
        doc.add_paragraph(note)
    selected = {"Yontem": tables["Yontem"],
        "Maddeler": tables["Maddeler"].reindex(columns=["Madde", "p", "q (BH)", "Karar", "İşaretli etki (puan)"]),
        "Parametreler": tables["Parametreler"],
        "Optimizasyon": tables["Optimizasyon"][["Madde", "Hipotez", "Yakınsama", "Skor/kişi", "Quadrature doğrulandı"]]}
    for title, frame in selected.items():
        doc.add_heading(title, 1)
        table = doc.add_table(rows=1, cols=len(frame.columns)); table.style = "Light Shading Accent 1"
        for cell, name in zip(table.rows[0].cells, frame.columns): cell.text = str(name)
        for values in frame.itertuples(index=False, name=None):
            for cell, value in zip(table.add_row().cells, values):
                cell.text = f"{value:.5g}" if isinstance(value, (float, np.floating)) else str(value)
    doc.add_paragraph("Tüm ayrıntılı tablolar, noktasal eğri aralıkları ve betimsel kategori tanıları XLSX/CSV eklerindedir. "
                      "Pozitif işaretli etki aynı theta'da odak grup lehine daha yüksek beklenen puandır; "
                      "ortalama etki eğrilerdeki yön değişimini gizleyebilir.")
    for figure in figures: doc.add_picture(io.BytesIO(figure), width=Inches(6))
    doc.add_heading("Kaynaklar", 1)
    for reference in REFERENCES: doc.add_paragraph(reference)
    buffer = io.BytesIO(); doc.save(buffer); return buffer.getvalue()


def bundle(out, tables=None):
    tables = report_tables(out) if tables is None else tables
    figures = []
    for name in out["table"].loc[~out["table"]["Sabit çapa"], "Madde"]:
        category = tables["Kategori_egrileri"].loc[lambda f: f.Madde == name]
        score = tables["Puan_farki"].loc[lambda f: f.Madde == name]
        figures.append(png_bytes(category, score, name, out["reference"], out["focal"]))
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("GRM_DIF.docx", word_bytes(out, tables, figures))
        archive.writestr("GRM_DIF.xlsx", export_xlsx(tables))
        for name, frame in tables.items():
            archive.writestr(f"{name}.csv", safe_frame(frame).to_csv(index=False).encode("utf-8-sig"))
        for i, figure in enumerate(figures): archive.writestr(f"item_{i + 1}_300dpi.png", figure)
        archive.writestr("README.txt", (out["status"] + "\n\n" + "\n\n".join(CAUTIONS + out["warnings"] + REFERENCES)
            + "\n\nKişi kimlikleri ve ham yanıtlar dahil değildir. Tüm tablo CSV dosyaları UTF-8 BOM içerir. "
              "Hesaplamalar rapor için yeniden kestirilmez. Grafik dosyaları 300 dpi. "
              "Samejima metadata: ETS RB-68-02; Chalmers metadata: yazarın mirt CITATION kaydı (2026-10-23 doğrulandı).\n").encode("utf-8"))
    return buffer.getvalue()
