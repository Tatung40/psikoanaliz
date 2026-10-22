"""CTT Word report with evidence and actionable but non-automatic item review."""
import io
from docx import Document
from core.ctt_dual_engine import item_guidance

def ctt_docx(result):
    doc=Document()
    doc.add_heading("EserPsychometrics - KTK Raporu",0)
    doc.add_paragraph(f"N={result['N']}; madde={result['k']}; listwise excluded={result['removed']}.")
    doc.add_paragraph(f"Cronbach alpha (native): {result['native_alpha']:.4f}")
    if result["reference_alpha"] is not None:
        doc.add_paragraph(f"Cronbach alpha (Pingouin): {result['reference_alpha']:.4f}. "
            f"Absolute difference: {result['difference']:.7f}; CI: {result['ci']}")
    else:
        doc.add_paragraph("Pingouin unavailable: cross-check not performed.")
    doc.add_heading("Madde inceleme tablosu",1)
    frame=item_guidance(result)
    table=doc.add_table(rows=1,cols=4)
    table.style="Table Grid"
    for c,v in zip(table.rows[0].cells,["Madde","p","r_ir","Tarama"]):c.text=v
    for _,r in frame.iterrows():
        row=table.add_row().cells
        for c,v in zip(row,[str(r["Madde"]),str(round(r["p"],3)),
                            str(round(r["r_ir"],3)),str(r["Durum"])]):c.text=v
    doc.add_heading("Bağımsız madde hesaplama denetimi",1)
    flagged=result["item_crosscheck"]
    doc.add_paragraph(f"{len(flagged)} madde karşılaştırıldı; {(flagged['Karar']=='YÖNTEM/FARK İNCELENMELİ').sum()} maddede hesaplama farkı gözlendi.")
    doc.add_heading("Karar ve bilimsel sınırlar",1)
    doc.add_paragraph("Düşük düzeltilmiş madde-toplam ilişkisi ya da uç madde güçlüğü "
        "otomatik silme nedeni değildir. İçerik, anahtar, hedef grup ve çeldiriciler "
        "uzmanlar tarafından değerlendirilmelidir. Cronbach alfa geçerliği veya tek boyutluluğu kanıtlamaz.")
    doc.add_heading("Kaynaklar",1)
    doc.add_paragraph("Cronbach, L. J. (1951). Coefficient alpha and the internal structure of tests. Psychometrika, 16, 297-334.")
    doc.add_paragraph("Sijtsma, K. (2009). On the use, the misuse, and the very limited usefulness of Cronbach's alpha. Psychometrika, 74, 107-120.")
    out=io.BytesIO();doc.save(out);return out.getvalue()
