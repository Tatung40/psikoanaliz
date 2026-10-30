"""Offline, deterministic Word report from an already validated academic result.

No API, RAG, numerical re-estimation, or fabricated interpretation.
"""
from io import BytesIO

from docx import Document
from docx.shared import Inches
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

from core.academic_result_schema import validate_result
from core.academic_visual_catalog import build_visual_catalog


def _value(value):
    if value is None:
        return "Raporlanmadı"
    if isinstance(value, bool):
        return "Evet" if value else "Hayır"
    if isinstance(value, float):
        return format(value, ".6g")
    return str(value)


def _figure(rows, spec):
    """Plot only real numeric observations; never interpolate missing statistics."""
    x, y = spec["x"], spec["y"]
    points = [(r[x], r[y]) for r in rows
              if isinstance(r.get(x), (int, float))
              and not isinstance(r.get(x), bool)
              and isinstance(r.get(y), (int, float))
              and not isinstance(r.get(y), bool)]
    if not points:
        return None
    fig, ax = plt.subplots(figsize=(6.4, 3.6))
    ax.plot([p[0] for p in points], [p[1] for p in points], linewidth=1.3)
    ax.set_xlabel(x)
    ax.set_ylabel(y)
    ax.grid(alpha=0.2)
    fig.tight_layout()
    stream = BytesIO()
    fig.savefig(stream, format="png", dpi=160)
    plt.close(fig)
    stream.seek(0)
    return stream


def build_academic_docx(result, *, max_rows_per_table=100):
    """Return DOCX bytes; report all omitted rows explicitly, no external requests."""
    source = validate_result(result)
    if type(max_rows_per_table) is not int or max_rows_per_table < 1:
        raise ValueError("max_rows_per_table must be positive")
    manifest = build_visual_catalog(source)
    doc = Document()
    doc.add_heading("Bulgular ve Yorumlar — Sayısal Rapor", 0)
    doc.add_paragraph(f"Alan: {source['family']} | Model: {source['model']}")
    doc.add_paragraph(f"N = {source['sample']['n']} | Madde sayısı: {_value(source['sample']['k'])}")
    doc.add_paragraph("Bu belge hesaplanan sonuçları sunar; kaynak temelli yorum ve istatistiksel çıkarım üretmez.")
    if source["fit_statistics"]:
        doc.add_heading("Model uyumu", level=1)
        for key, value in source["fit_statistics"].items():
            if isinstance(value, (list, dict)):
                doc.add_paragraph(f"{key}: yapılandırılmış veri; ilgili tablolara bakınız.")
            else:
                doc.add_paragraph(f"{key}: {_value(value)}")
    doc.add_heading("Tablolar", level=1)
    for spec in manifest["tables"]:
        rows = source["tables"][spec["source_table"]]
        doc.add_paragraph(spec["apa_label"], style="Normal")
        doc.add_paragraph(spec["title"])
        if not rows:
            doc.add_paragraph("Tabloda gözlem bulunmuyor.")
            continue
        columns = list(dict.fromkeys(k for row in rows for k in row))
        table = doc.add_table(rows=1, cols=len(columns))
        for cell, name in zip(table.rows[0].cells, columns):
            cell.text = str(name)
        for row in rows[:max_rows_per_table]:
            cells = table.add_row().cells
            for cell, name in zip(cells, columns):
                cell.text = _value(row.get(name))
        if len(rows) > max_rows_per_table:
            doc.add_paragraph(f"Not. {len(rows)-max_rows_per_table} satır belge boyutu nedeniyle gösterilmedi; özgün JSON'da korunur.")
    doc.add_heading("Şekiller", level=1)
    for spec in manifest["figures"]:
        plot = _figure(source["tables"][spec["source_table"]], spec)
        if plot is None:
            doc.add_paragraph(f"{spec['apa_label']}: sayısal grafik noktaları yok; çizilmedi.")
            continue
        doc.add_paragraph(spec["apa_label"])
        doc.add_paragraph(spec["title"])
        doc.add_picture(plot, width=Inches(5.8))
        doc.add_paragraph(f"Not. Kaynak: {spec['source_table']}; x={spec['x']}, y={spec['y']}.")
    warnings = list(source["warnings"]) + manifest["warnings"]
    if warnings:
        doc.add_heading("Bilimsel sınırlılıklar ve uyarılar", level=1)
        for warning in warnings:
            doc.add_paragraph(str(warning), style="List Bullet")
    doc.add_heading("Tekrarlanabilirlik bilgisi", level=1)
    for key, value in source["provenance"].items():
        doc.add_paragraph(f"{key}: {_value(value)}")
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()
