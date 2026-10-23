import io
import zipfile
from pathlib import Path
import numpy as np
import pytest
from PIL import Image
from docx import Document
from openpyxl import load_workbook
from core import polytomous_dif as P
from reports.polytomous_dif_report import report_tables, bundle
from scripts.polytomous_dif_validation import generate, ANCHORS
from ui.polytomous_dif_view import result_key


@pytest.mark.parametrize("model", P.MODELS)
def test_qualified_bundle_safe_labels_no_rows_no_refitting(model, monkeypatch):
    x, g, _ = generate(model)
    x = x.rename(columns={"M5": "=unsafe"})
    labels = None
    if model == "NRM":
        labels = ["Z", "=A", "K"]; x = x.replace(dict(enumerate(labels)))
    out = P.test_dif(x, g, "R", "F", ANCHORS, model=model, category_labels=labels)
    out["private_person_ids"] = ["PRIVATE_PERSON_ABC"]
    tables = report_tables(out)
    assert tables["Katsayilar"].shape[0] == 36
    assert tables["Optimizasyon"]["Bilgi matrisi"].all()
    assert tables["Kategori_tanilari"].groupby(["Grup", "Madde"])["Model oranı"].sum().between(.999999, 1.000001).all()
    assert ("Puan_farki" in tables) == (model != "NRM")
    if model == "NRM":
        assert set(tables["Kategori_egrileri"].Kategori) == set(labels)
        assert all("puan" not in e.lower() for e in tables["Etkiler"].Etki)
        assert tables["Parametreler"].Parametre.str.contains("/ Z").all()
        assert tables["Etkiler"][["%95 GA alt", "%95 GA üst"]].isna().all().all()
        assert tables["Etkiler"]["Aralık durumu"].str.contains("%88").all()
    monkeypatch.setattr(P.Fitter, "fit", lambda *a, **k: pytest.fail("Report refitted a model"))
    data = bundle(out, tables)
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        assert {f"{model}_DIF.docx", f"{model}_DIF.xlsx", "Etkiler.csv", "README.txt", "item_1_300dpi.png"} <= set(z.namelist())
        for name in z.namelist():
            assert ".." not in name and not name.startswith("/")
            if name.endswith(".csv"):
                assert z.read(name).startswith(b"\xef\xbb\xbf")
                assert b"PRIVATE_PERSON_ABC" not in z.read(name)
        assert "'=unsafe" in z.read("Maddeler.csv").decode("utf-8-sig")
        wb = load_workbook(io.BytesIO(z.read(f"{model}_DIF.xlsx")))
        cells = [cell for row in wb["Maddeler"] for cell in row if "unsafe" in str(cell.value)]
        assert cells and all(c.data_type == "s" for c in cells)
        doc = Document(io.BytesIO(z.read(f"{model}_DIF.docx")))
        text = " ".join(p.text for p in doc.paragraphs)
        assert "EXPERIMENTAL" in text and "R mirt bu çalışmada çalıştırılmadı" in text
        assert "PRIVATE_PERSON_ABC" not in text
        image = Image.open(io.BytesIO(z.read("item_1_300dpi.png")))
        assert image.size == (2100, 2400) and image.info["dpi"][0] == pytest.approx(300, abs=.1)


def test_fingerprint_tracks_nominal_labels_model_settings_not_ids():
    x, groups, _ = generate()
    x["group"] = groups; x["ID"] = np.arange(len(x))
    items = list(x.columns[:6])
    args = (items, "group", "R", "F", ANCHORS, "NRM", 3, [0, 1, 2], 61)
    key = result_key(x, *args)
    altered = x.copy(); altered.ID += 1000
    assert result_key(altered, *args) == key
    altered.iloc[0, 0] = (altered.iloc[0, 0] + 1) % 3
    assert result_key(altered, *args) != key
    assert result_key(x, items, "group", "R", "F", ANCHORS, "NRM", 3, [2, 0, 1], 61) != key
    assert result_key(x, items, "group", "R", "F", ANCHORS, "PCM", 3, None, 61) != key


def test_app_model_switch_nominal_labels_stale_results_and_invalid_data():
    from streamlit.testing.v1 import AppTest
    x, groups, _ = generate("PCM", n_reference=300, n_focal=300)
    x["Grup"] = groups; x["ID"] = np.arange(len(x))
    at = AppTest.from_file(str(Path(__file__).resolve().parents[1] / "app.py"), default_timeout=180)
    for key, value in dict(raw_df=x, scored_df=x.copy(), item_cols=list(x.columns), id_col="ID", group_col="Grup",
                           is_dichotomous=False, nav_page="MTK: Çok Kategorili Formel DIF").items(): at.session_state[key] = value
    at.run()
    assert not at.exception and at.button(key="poly_run").disabled
    assert set(at.multiselect(key="poly_items").value) == set(x.columns[:6])
    assert list(at.selectbox(key="poly_model").options) == list(P.MODELS)
    at.multiselect(key="poly_anchors").set_value(ANCHORS).run()
    at.checkbox(key="poly_ack").check().run()
    at.button(key="poly_run").click().run()
    assert not at.exception and at.session_state["poly_dif_result"]["out"]["model"] == "PCM"
    at.selectbox(key="poly_model").set_value("GPCM").run()
    assert any("eski sonuç gösterilmiyor" in info.value for info in at.info)
    at.selectbox(key="poly_model").set_value("RSM").run()
    assert any("yalnız madde konumu" in caption.value for caption in at.caption)
    at.selectbox(key="poly_model").set_value("NRM").run()
    at.text_input(key="poly_labels").set_value('not JSON').run()
    assert at.button(key="poly_run").disabled
    assert not at.exception
    nominal, ng, _ = generate("NRM", n_reference=300, n_focal=300)
    nominal = nominal.replace({0: "Z", 1: "A", 2: "K"}); nominal["Grup"] = ng; nominal["ID"] = np.arange(len(nominal))
    at.session_state["raw_df"] = nominal
    at.text_input(key="poly_labels").set_value('["Z","A","K"]').run()
    at.button(key="poly_run").click().run()
    assert not at.exception and at.session_state["poly_dif_result"]["out"]["model"] == "NRM"
    at.text_input(key="poly_labels").set_value('["Z","A","UNKNOWN"]').run()
    at.button(key="poly_run").click().run()
    assert not at.exception and any("etiket" in error.value for error in at.error)
    assert "poly_dif_result" not in at.session_state
