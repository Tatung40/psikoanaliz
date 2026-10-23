import io
import zipfile
from pathlib import Path
import numpy as np
import pandas as pd
import pytest
from docx import Document
from openpyxl import load_workbook
from PIL import Image
from core import grm_dif as G
from reports.grm_dif_report import report_tables, bundle
from scripts.grm_dif_validation import generate
from ui.grm_dif_view import result_key


def test_research_bundle_has_safe_tables_curves_uncertainty_and_no_responses():
    x, groups, _ = generate(n_reference=300, n_focal=300)
    x = x.rename(columns={"M5": "=dangerous_label"})
    out = G.test_dif(x, groups, "R", "F", list(x.columns[:4]), points=61)
    out["private_person_ids"] = ["PRIVATE_PERSON_123"]
    tables = report_tables(out)
    assert tables["Katsayilar"].shape[0] == 12
    assert tables["Optimizasyon"]["Bilgi matrisi"].all()
    assert tables["Kategori_tanilari"].groupby(["Grup", "Madde"])["Model oranı"].sum().between(.999999, 1.000001).all()
    data = bundle(out, tables)
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        assert {"GRM_DIF.docx", "GRM_DIF.xlsx", "README.txt", "item_1_300dpi.png", "item_2_300dpi.png"} <= set(z.namelist())
        for name in z.namelist():
            assert ".." not in name and not name.startswith("/")
            if name.endswith(".csv"):
                raw = z.read(name)
                assert raw.startswith(b"\xef\xbb\xbf") and b"PRIVATE_PERSON_123" not in raw
        csv = z.read("Maddeler.csv").decode("utf-8-sig")
        assert "'=dangerous_label" in csv
        wb = load_workbook(io.BytesIO(z.read("GRM_DIF.xlsx")))
        cells = [cell for row in wb["Maddeler"] for cell in row if "dangerous_label" in str(cell.value)]
        assert cells and all(cell.data_type == "s" for cell in cells)
        doc = Document(io.BytesIO(z.read("GRM_DIF.docx")))
        text = " ".join(p.text for p in doc.paragraphs)
        assert "R mirt bu doğrulamada çalıştırılmadı" in text
        assert "EXPERIMENTAL" in text and "PRIVATE_PERSON_123" not in text
        for name in ("item_1_300dpi.png", "item_2_300dpi.png"):
            picture = Image.open(io.BytesIO(z.read(name)))
            assert picture.size == (2100, 2400)
            assert picture.info["dpi"][0] == pytest.approx(300, abs=.1)


def test_result_fingerprint_tracks_data_settings_but_not_private_ids():
    x, group, _ = generate(n_reference=100, n_focal=100)
    x["group"] = group; x["ID"] = np.arange(len(x))
    items, anchors = list(x.columns[:6]), list(x.columns[:4])
    args = (items, "group", "R", "F", anchors, 3, 61)
    key = result_key(x, *args)
    altered = x.copy(); altered["ID"] += 1000
    assert result_key(altered, *args) == key
    altered.loc[0, "M1"] = (altered.loc[0, "M1"] + 1) % 3
    assert result_key(altered, *args) != key
    assert result_key(x, items, "group", "R", "F", anchors[:-1], 3, 61) != key
    assert result_key(x, items, "group", "F", "R", anchors, 3, 61) != key


def test_app_separates_experimental_grm_from_screens_and_invalidates_stale_data():
    from streamlit.testing.v1 import AppTest
    app = str(Path(__file__).resolve().parents[1] / "app.py")
    x, groups, _ = generate(n_reference=300, n_focal=300)
    x["Grup"] = groups; x["OGRENCI_NO"] = np.arange(len(x))
    at = AppTest.from_file(app, default_timeout=180)
    for key, value in {"raw_df": x, "scored_df": x.copy(), "item_cols": list(x.columns),
                       "id_col": "OGRENCI_NO", "group_col": "Grup", "is_dichotomous": False,
                       "nav_page": "MTK: GRM Formel DIF"}.items():
        at.session_state[key] = value
    at.run()
    assert not at.exception and at.button(key="grm_run").disabled
    assert set(at.multiselect(key="grm_items").value) == set(x.columns[:6])
    at.multiselect(key="grm_anchors").set_value(list(x.columns[:4])).run()
    at.checkbox(key="grm_ack").check().run()
    at.button(key="grm_run").click().run()
    assert not at.exception
    out = at.session_state["grm_dif_result"]["out"]
    assert out["inputs"]["names"] == list(x.columns[:6])
    assert "EXPERIMENTAL" in out["status"]
    at.multiselect(key="grm_anchors").set_value(list(x.columns[:3])).run()
    assert any("eski sonuç gösterilmiyor" in info.value for info in at.info)
    bad = x.copy()
    for col in x.columns[:6]:
        bad[col] = bad[col].astype(str).replace({"0": "A", "1": "B", "2": "C"})
    at.session_state["raw_df"] = bad
    at.button(key="grm_run").click().run()
    assert not at.exception and any("nominal" in error.value for error in at.error)
