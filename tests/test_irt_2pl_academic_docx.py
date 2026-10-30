from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document
import pandas as pd
import pytest

from reports.irt_2pl_academic_docx import export_2pl_docx


def _report():
    return SimpleNamespace(
        model="2PL",
        metadata={"N": 80, "Madde": 2, "Yakınsama": True, "Motor": "Python mirt",
                  "Kestirim": "ML", "AIC": 120.0},
        parameters=pd.DataFrame({"Madde": ["M1", "M2"], "a": [1.2, 0.8],
                                 "b": [-0.3, 0.4]}),
        curves=pd.DataFrame([
            {"Madde": item, "Theta": theta, "Kategori": str(cat),
             "Olasılık": p if cat else 1-p}
            for item in ("M1", "M2")
            for theta, p in [(-1., 0.2), (0., 0.5), (1., 0.8)]
            for cat in (0, 1)]),
        information=pd.DataFrame({"θ": [-1., 0., 1.],
                                  "I_M1": [0.2, 0.3, 0.2],
                                  "I_M2": [0.1, 0.2, 0.1],
                                  "Test bilgisi": [0.3, 0.5, 0.3]}),
        errors=[], warnings=[])


def test_2pl_word_groups_each_item_and_information():
    with patch("reports.irt_2pl_academic_docx.build_report", return_value=_report()):
        raw = export_2pl_docx(object(), ["M1", "M2"])
    doc = Document(BytesIO(raw))
    assert len(doc.inline_shapes) == 3
    assert any("M1 ICC" in p.text for p in doc.paragraphs)
    assert any("M2 ICC" in p.text for p in doc.paragraphs)
    assert any("120" in p.text for p in doc.paragraphs)


def test_2pl_word_rejects_unconverged_result():
    report = _report()
    report.metadata["Yakınsama"] = False
    with patch("reports.irt_2pl_academic_docx.build_report", return_value=report):
        with pytest.raises(ValueError, match="Convergence"):
            export_2pl_docx(object(), ["M1", "M2"])
