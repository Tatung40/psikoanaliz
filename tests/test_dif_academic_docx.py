from io import BytesIO
from unittest.mock import patch

import pandas as pd
import pytest
from docx import Document

from reports.dif_academic_docx import dif_academic_docx


def _result():
    return {"metadata": {"Model": "2PL", "N_reference": 50, "N_focal": 50,
                          "Items": ["M1"], "Inference": False,
                          "Correction_applied": False, "Status": "NOT CERTIFIED",
                          "Limitations": "NOT CERTIFIED"},
            "tables": {"Parameters": pd.DataFrame({"Item": ["M1"], "Value": [0.75]}),
                       "DIF": pd.DataFrame({"Item": ["M1"], "LR": [1.5],
                                            "p": [float("nan")]})}}


def test_dif_word_contains_source_tables_and_four_panel_image():
    import matplotlib.pyplot as plt
    fig, ax = plt.subplots()
    ax.plot([0, 1], [0.2, 0.8])
    stream = BytesIO()
    fig.savefig(stream, format="png")
    plt.close(fig)
    with patch("reports.dif_academic_docx.png", return_value=stream.getvalue()):
        raw = dif_academic_docx(_result(), "M1")
    doc = Document(BytesIO(raw))
    assert len(doc.inline_shapes) == 1
    assert any("NOT CERTIFIED" in p.text for p in doc.paragraphs)
    assert any("M1" in cell.text for table in doc.tables for row in table.rows for cell in row.cells)


def test_dif_word_rejects_unknown_item():
    with pytest.raises(ValueError, match="absent"):
        dif_academic_docx(_result(), "M2")
