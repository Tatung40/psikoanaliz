from io import BytesIO

from docx import Document

from core.academic_result_schema import build_result
from reports.academic_docx import build_academic_docx


def test_offline_docx_preserves_observed_values_and_embeds_figure():
    result = build_result(family="IRT", model="2PL", n=100, k=2,
        tables={"category_curves": [
            {"Theta": -1.0, "Olasılık": 0.2},
            {"Theta": 0.0, "Olasılık": 0.5},
            {"Theta": 1.0, "Olasılık": 0.8},
        ]}, fit_statistics={"AIC": 123.456},
        metadata={"analysis_type": "item_calibration"},
        warnings=["Inference not certified"])
    raw = build_academic_docx(result)
    doc = Document(BytesIO(raw))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "123.456" in text
    assert "Inference not certified" in text
    assert "Şekil 1" in text
    assert len(doc.inline_shapes) == 1
    assert doc.tables[0].rows[1].cells[1].text == "0.2"


def test_missing_graph_is_not_fabricated():
    result = build_result(family="IRT", model="2PL", n=10,
        tables={"category_curves": [{"Theta": None, "Olasılık": None}]},
        metadata={"analysis_type": "item_calibration"})
    doc = Document(BytesIO(build_academic_docx(result)))
    assert len(doc.inline_shapes) == 0
