"""The 'Python mirt ile hesapla' 2PL page (the screen users and the sample-data loader reach) ends with
the academic Word download; the DOCX opens and contains the fitted tables and ICC/information figures."""
import io
from pathlib import Path

import pytest

pytest.importorskip("mirt")
from docx import Document
from streamlit.testing.v1 import AppTest

from core import example_datasets as E

ROOT = Path(__file__).resolve().parents[1]


def _app(model, case):
    b = E.materialize(case); x = b["data"][b["items"]].astype(int)
    at = AppTest.from_file(str(ROOT / "app.py"), default_timeout=300)
    for key, value in dict(raw_df=x, scored_df=x.copy(), item_cols=list(x.columns), id_col=None, group_col=None,
                           is_dichotomous=True, nav_page=f"MTK: Python mirt {model}").items():
        at.session_state[key] = value
    at.run()
    at.button(key=f"python_mirt_start_{model}").click().run()
    assert not at.exception, [e.value for e in at.exception]
    return at, list(x.columns)


def _downloads(at):
    """Downloads of the analysis page itself (the shared Python Kodları panel renders after every page)."""
    return [d for d in at.get("download_button") if d.proto.id and not any(k in d.proto.id for k in ("repro_py", "repro_zip"))]


def test_2pl_word_download_is_last_output_and_docx_holds_real_results():
    at, items = _app("2PL", "irt_006")
    labels = [d.proto.label for d in _downloads(at)]
    assert labels[-1] == "📄 Akademik 2PL Raporu — Word İndir", labels
    assert any(m.value == "### 📄 Akademik 2PL Raporu" for m in at.markdown)
    store = at.session_state["python_mirt_2pl_word"]
    assert store["error"] is None and store["key"] == at.session_state["python_mirt_fit_2PL"]["key"]
    doc = Document(io.BytesIO(store["docx"]))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert len(doc.tables) >= 1 and len(doc.inline_shapes) >= len(items) + 1
    assert all(t.style.name == 'Table Grid' for t in doc.tables)
    assert "Madde karakteristik eğrileri (ICC)" in text and "Madde ve test bilgi fonksiyonları" in text
    cells = " ".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert all(item in cells or item in text for item in items)
    assert any('fit_python_mirt' in code.code for code in at.code), '2PL Python replay code absent'
    first = store["docx"]
    at.run()                                                    # ordinary rerun: button stays, no rebuild
    assert [d.proto.label for d in _downloads(at)][-1] == "📄 Akademik 2PL Raporu — Word İndir"
    assert at.session_state["python_mirt_2pl_word"]["docx"] is first


@pytest.mark.parametrize("model,case", [("1PL", "irt_003"), ("3PL", "irt_009")])
def test_other_models_have_no_2pl_word_button(model, case):
    at, _ = _app(model, case)
    assert "📄 Akademik 2PL Raporu — Word İndir" not in [d.proto.label for d in _downloads(at)]
